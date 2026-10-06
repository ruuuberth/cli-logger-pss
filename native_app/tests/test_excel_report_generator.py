"""Tests for ExcelReportGenerator multi-sheet support (H2H single workbook)."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.reporting.report_generator import ExcelReportGenerator, ReportConfig


def _generator(tmp_path, title="test report") -> ExcelReportGenerator:
    return ExcelReportGenerator(
        ReportConfig(
            title=title,
            output_path=tmp_path,
            include_timestamp=False,
            format="excel",
        )
    )


def test_multi_sheet_workbook_sheetnames_and_order(tmp_path) -> None:
    gen = _generator(tmp_path)
    gen.add_sheet("Resumen", [{"Métrica": "Batallas", "Valor": 10}])
    gen.add_sheet("Batallas", [{"Battle ID": "b1", "Resultado": "VICTORY"}])
    gen.add_sheet("Tendencias", [{"Periodo": "2024-01", "Win Rate": "80%"}])
    fleet = [{"Jugador": "Alice", "Tipo": "Nave", "Nombre": "Acorazado"}]
    gen.add_sheet("Flota", fleet)
    output_file = gen.generate()

    assert output_file.exists()
    assert output_file.name == "test_report.xlsx"

    wb = load_workbook(output_file)
    assert wb.sheetnames == ["Resumen", "Batallas", "Tendencias", "Flota"]
    assert wb.active is not None and wb.active.title == "Resumen"


def test_multi_sheet_headers_values_and_styling(tmp_path) -> None:
    gen = _generator(tmp_path)
    gen.add_sheet("Resumen", [{"Métrica": "Total Batallas", "Valor": 7}])
    gen.add_sheet("Batallas", [{"Battle ID": "b1", "Resultado": "VICTORY"}])
    output_file = gen.generate()

    wb = load_workbook(output_file)

    ws_resumen = wb["Resumen"]
    assert ws_resumen.cell(row=1, column=1).value == "Métrica"
    assert ws_resumen.cell(row=1, column=2).value == "Valor"
    assert ws_resumen.cell(row=2, column=1).value == "Total Batallas"
    assert ws_resumen.cell(row=2, column=2).value == 7
    header_cell = ws_resumen.cell(row=1, column=1)
    assert header_cell.fill.start_color.rgb == "004472C4"
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb == "00FFFFFF"
    assert header_cell.alignment.horizontal == "center"
    data_cell = ws_resumen.cell(row=2, column=1)
    assert data_cell.alignment.horizontal == "left"

    ws_batallas = wb["Batallas"]
    assert ws_batallas.cell(row=1, column=1).value == "Battle ID"
    assert ws_batallas.cell(row=2, column=2).value == "VICTORY"


def test_multi_sheet_column_widths(tmp_path) -> None:
    gen = _generator(tmp_path)
    gen.add_sheet("Resumen", [{"Métrica": "Total Batallas", "Valor": 7}])
    output_file = gen.generate()

    wb = load_workbook(output_file)
    ws = wb["Resumen"]
    # max(len("Métrica")=7, len("Total Batallas")=14) + 2 = 16, capped at 50
    assert ws.column_dimensions["A"].width == 16
    assert ws.column_dimensions["B"].width == 7


def test_add_rows_routes_to_default_data_sheet(tmp_path) -> None:
    """Battle-report path: add_rows() produces exactly one 'Data' sheet."""
    gen = _generator(tmp_path)
    gen.add_rows([{"Battle ID": "b1", "Resultado": "VICTORY"}, {"Battle ID": "b2", "Resultado": "DEFEAT"}])
    output_file = gen.generate()

    wb = load_workbook(output_file)
    assert wb.sheetnames == ["Data"]
    ws = wb["Data"]
    assert ws.cell(row=1, column=1).value == "Battle ID"
    assert ws.cell(row=2, column=1).value == "b1"
    assert ws.cell(row=3, column=2).value == "DEFEAT"


def test_add_rows_and_add_sheet_coexist(tmp_path) -> None:
    """add_rows targets the default sheet only; other sheets are separate."""
    gen = _generator(tmp_path)
    gen.add_rows([{"Col": "from add_rows"}])
    gen.add_sheet("Extra", [{"Col": "from add_sheet"}])
    output_file = gen.generate()

    wb = load_workbook(output_file)
    assert wb.sheetnames == ["Data", "Extra"]
    assert wb["Data"].cell(row=2, column=1).value == "from add_rows"
    assert wb["Extra"].cell(row=2, column=1).value == "from add_sheet"


def test_all_empty_sheets_raise(tmp_path) -> None:
    gen = _generator(tmp_path)
    with pytest.raises(ValueError, match="No hay datos para reportar"):
        gen.generate()


def test_single_empty_sheet_among_nonempty_is_allowed(tmp_path) -> None:
    """Per-sheet emptiness is the caller's choice; only all-empty raises."""
    gen = _generator(tmp_path)
    gen.add_sheet("Resumen", [{"Métrica": "x", "Valor": 1}])
    gen.add_sheet("Tendencias", [])
    output_file = gen.generate()

    wb = load_workbook(output_file)
    assert wb.sheetnames == ["Resumen", "Tendencias"]
    assert wb["Resumen"].cell(row=2, column=1).value == "x"
    assert wb["Tendencias"].max_row == 1  # empty sheet, nothing written


def test_empty_rows_list_via_add_rows_raises_on_generate(tmp_path) -> None:
    """All-empty rule applies to the default-sheet path too."""
    gen = _generator(tmp_path)
    gen.add_rows([])
    with pytest.raises(ValueError, match="No hay datos para reportar"):
        gen.generate()


def test_duplicate_sheet_name_raises(tmp_path) -> None:
    gen = _generator(tmp_path)
    gen.add_sheet("Batallas", [{"A": 1}])
    with pytest.raises(ValueError, match="Batallas"):
        gen.add_sheet("Batallas", [{"A": 2}])
    # default sheet colliding with a named sheet must also raise
    gen2 = _generator(tmp_path)
    gen2.add_sheet("Data", [{"A": 1}])
    with pytest.raises(ValueError, match="Data"):
        gen2.add_rows([{"A": 2}])


def test_too_long_sheet_name_raises(tmp_path) -> None:
    gen = _generator(tmp_path)
    with pytest.raises(ValueError, match="demasiado largo"):
        gen.add_sheet("x" * 32, [{"A": 1}])


def test_invalid_sheet_name_characters_raise(tmp_path) -> None:
    gen = _generator(tmp_path)
    for bad in ["My[Sheet]", "My]Sheet", "My:Sheet", "My*Sheet", "My?Sheet", "My/Sheet", "My\\Sheet"]:
        with pytest.raises(ValueError, match="no válidos"):
            gen.add_sheet(bad, [{"A": 1}])


def test_more_than_26_columns_saves_successfully(tmp_path) -> None:
    """Regression for #22: >26 columns used to write invalid column letters."""
    rows = [{f"col_{i}": f"value_{i}" for i in range(1, 31)}]
    gen = _generator(tmp_path)
    gen.add_rows(rows)
    output_file = gen.generate()  # must not raise on save

    wb = load_workbook(output_file)
    ws = wb["Data"]
    assert ws.cell(row=1, column=1).value == "col_1"
    assert ws.cell(row=1, column=27).value == "col_27"
    assert ws.cell(row=1, column=30).value == "col_30"
    assert ws.cell(row=2, column=30).value == "value_30"
    # column_dimensions keys are valid letters, including past Z
    assert ws.column_dimensions["AA"].width is not None
    assert ws.column_dimensions["AD"].width == min(len("value_30") + 2, 50)
    assert "[" not in {str(k) for k in ws.column_dimensions.keys()}


def test_empty_name_falls_back_to_default_sheet(tmp_path) -> None:
    gen = _generator(tmp_path)
    gen.add_sheet("", [{"A": 1}])
    assert list(gen._sheets.keys()) == ["Data"]
    output_file = gen.generate()
    wb = load_workbook(output_file)
    assert wb.sheetnames == ["Data"]


# --- GenerateH2HReportCommand Excel path (single multi-sheet workbook) ---


def _h2h_data(with_fleet: bool) -> dict:
    data: dict = {
        "summary": {
            "player_low_name": "Alice",
            "player_high_name": "Bob",
            "player_low_wins": 2,
            "player_high_wins": 1,
            "total_battles": 3,
        },
        "battles": [
            {"battle_id": "b1", "captured_at": "2024-01-01T10:00:00", "attacker_name": "Alice",
             "defender_name": "Bob", "outcome": "VICTORY", "attacker_trophy_delta": 10,
             "defender_trophy_delta": -10, "loot_minerals": 100, "loot_gas": 50},
        ],
        "trends": [
            {"period": "2024-01", "battle_count": 3, "player_low_wins": 2,
             "player_high_wins": 1, "player_low_avg_trophies": 4000, "player_high_avg_trophies": 3900},
        ],
    }
    if with_fleet:
        data["fleet_breakdown"] = {
            "battles_analyzed": 3,
            "player_low": {
                "user_id": 10,
                "total_battles_analyzed": 3,
                "ships": [{"id": 11, "battles": 3, "ship_name": "ShipA", "ship_design_name": "Acorazado",
                           "avg_level": 8, "avg_power_score": 4000.0}],
                "rooms": [],
                "crew": [],
            },
            "player_high": {"user_id": 20, "total_battles_analyzed": 3, "ships": [], "rooms": [], "crew": []},
        }
    else:
        data["fleet_breakdown"] = {"player_low": {}, "player_high": {}}
    return data


def _run_h2h_excel(tmp_path, monkeypatch, h2h_data: dict) -> tuple[int, Path]:
    """Drive GenerateH2HReportCommand.execute() non-interactively (Excel branch)."""
    import app.cli.concrete_commands as cc

    pair = {
        "low_user_id": 10,
        "high_user_id": 20,
        "low_name": "Alice",
        "high_name": "Bob",
    }

    class _FakeService:
        def get_unique_player_pairs(self):
            return [pair]

        def get_h2h_report_data(self, **kwargs):
            return h2h_data

    # Prompt answers in call order: selection method, pair index.
    # Any further prompt falls back to its default.
    answers = iter(["1", "0"])

    def _fake_prompt(prompt="", default="", **kwargs):
        try:
            return next(answers)
        except StopIteration:
            return default

    monkeypatch.setattr(cc, "prompt_input", _fake_prompt)
    monkeypatch.setattr(cc.ApiFlowCliService, "get_unique_player_pairs", _FakeService().get_unique_player_pairs)
    monkeypatch.setattr(cc.ApiFlowCliService, "get_h2h_report_data", _FakeService().get_h2h_report_data)

    cmd = cc.GenerateH2HReportCommand(runtime=None)
    out_dir = tmp_path / "reports"
    rc = cmd.execute([
        "--non-interactive",
        "--format=excel",
        f"--output-dir={out_dir}",
        "--filename=H2H_Alice_vs_Bob",
        "--no-timestamp",
    ])
    return rc, out_dir


def test_h2h_command_generates_single_multi_sheet_workbook(tmp_path, monkeypatch) -> None:
    rc, out_dir = _run_h2h_excel(tmp_path, monkeypatch, _h2h_data(with_fleet=True))

    assert rc == 0
    files = list(out_dir.glob("*.xlsx"))
    assert len(files) == 1, f"expected exactly one workbook, got {files}"

    from openpyxl import load_workbook
    wb = load_workbook(files[0])
    assert wb.sheetnames == ["Resumen", "Batallas", "Tendencias", "Flota"]
    assert wb["Resumen"].cell(row=2, column=1).value == "Jugador 1"
    assert wb["Batallas"].cell(row=2, column=2).value == "b1"
    assert wb["Tendencias"].cell(row=2, column=1).value == "2024-01"
    assert wb["Flota"].cell(row=2, column=3).value == "Acorazado"


def test_h2h_command_omits_flota_sheet_when_no_fleet_rows(tmp_path, monkeypatch) -> None:
    rc, out_dir = _run_h2h_excel(tmp_path, monkeypatch, _h2h_data(with_fleet=False))

    assert rc == 0
    files = list(out_dir.glob("*.xlsx"))
    assert len(files) == 1

    from openpyxl import load_workbook
    wb = load_workbook(files[0])
    assert wb.sheetnames == ["Resumen", "Batallas", "Tendencias"]
